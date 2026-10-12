'''
OWI database

Dependencies:
    Python >= 3.7
    OWI_secrets.py (optional, git-ignored; see class ftp_params)

Some code is copied from convert_cwi_to_pandas.py by Randal Barnes     

The purpose of module OWI_download_ftp is to download specified zip files from  
the Minnesota Geologic Survey ftp site, and unzip the files.  

Method RUN_download_cwi demonstrates usage.

Manual download instruction:

'''

import os, shutil
import getpass, shlex
import time
import datetime
from ftplib import FTP
import zipfile
import csv

import logging

from OWI_config import versioninfo

try:
    import OWI_secrets as S     # git-ignored; may be missing
except ImportError:
    S = None


def xlsx_to_csv(xlsx_file, csv_file=None):
    """
    Converts an .xlsx file to .csv format on Ubuntu or any OS with Python.

    Parameters:
        xlsx_file (str): Path to the .xlsx file.
        csv_file (str, optional): Path to save the .csv file. Defaults to the same name as the .xlsx file in the same directory.

    Returns:
        bool: True on success.

    Requires openpyxl (imported here so the module loads without it).
    """

    # print (f"Found xlsx files rather than csv.\n > Converting: {xlsx_file}\n > To: {csv_file}")
    # Check if the .xlsx file exists
    if not os.path.exists(xlsx_file):
        raise FileNotFoundError(f"The file '{xlsx_file}' does not exist.")
    if csv_file is None:
        csv_file = os.path.splitext(xlsx_file)[0] + '.csv'

    # Read the Excel file and save it as CSV
    try:
        from openpyxl import load_workbook

        # Read the first sheet of the Excel file, row by row, into the CSV file
        wb = load_workbook(xlsx_file, read_only=True, data_only=True)
        try:
            with open(csv_file, 'w', newline='') as f:
                w = csv.writer(f)
                for row in wb.worksheets[0].iter_rows(values_only=True):
                    w.writerow(['' if v is None else v for v in row])
        finally:
            wb.close()
        assert os.path.exists(csv_file)
        #print(f" Converted '{xlsx_file}' to '{csv_file}'.")
    except Exception as e:
        raise RuntimeError(f"Failed to convert {xlsx_file} to CSV: {e}")
    
    return True



log = logging.getLogger('CWI_refresh')
logging.basicConfig(level=logging.INFO,
                    format = '%(asctime)s - %(name)s - %(levelname)s - %(message)s',
                    datefmt='%Y-%m-%d %H:%M:%S')

__author__ = "William C Olsen"
__version__ = "2021-10-10"

def download_cwi_from_ftp( ftpaddr,
                           ftppath,
                           ftpuser, 
                           ftppass,
                           downloadpath,
                           downloadshppath,
                           downloadfiles):
    ''' 
    Download newer cwi files from ftp. Unzip. Return list of new files. 
    
    This function downloads and extracts cwi files from the MGS ftp site.
    File dates are compared, and the ftp download only progresses if the
    ftp version is newer.  
    
    The ftp login is found by class ftp_params.
    The download targets are defined in OWI_config.py
    
    Notes on variable names:
        src is used for ftp source files
        dst is used for local copies of the files
    '''

    log.info(msg=('Downloading CWI sources from ftp site'))
    rv = []
    ftpstart = time.time()
    with FTP(ftpaddr) as ftp:
        ftp.login(user=ftpuser, passwd=ftppass)
        log.debug(msg='ftp Log-in successful')
        
        # Get the modify dates of the files on the ftp site. Formatted as integer, e.g. 20170521081548
        dict_ftp_file_modify_times = {f[0].lower():int(float(f[1].get('modify'))) for f in ftp.mlsd(path=ftppath)}
        for key,val in dict_ftp_file_modify_times.items():
            print (val, key)
    
        os.makedirs(downloadpath, exist_ok=True)
        os.makedirs(downloadshppath, exist_ok=True)
        
        for fname in downloadfiles:
            key = fname.lower()
            if 'locs' in fname: 
                destfile = os.path.join(downloadshppath, fname)
            else:
                destfile = os.path.join(downloadpath, fname)
            if os.path.exists(destfile) and key in dict_ftp_file_modify_times:
                # Compare file times to see if a fresh download is needed
                mtime = datetime.datetime.fromtimestamp( os.path.getmtime(destfile) )
                desttime = int(mtime.strftime('%Y%m%d%H%M%S'))
                
                if dict_ftp_file_modify_times[key] < desttime:
                    log.info(f'{fname} is up to date. Download skipped.')
                    continue
                
            start = time.time()
            print( f"downloading {fname} ...", end='' )
            
            with open(destfile, 'wb') as localfile:
                try:
                    ftp.retrbinary('RETR '+ftppath+'/'+fname, localfile.write, 1024 )
                except Exception as e:
                    lines = ftp.retrlines('LIST')
                    print (lines)
                    print (f"ERROR downloading {ftppath+'/'+fname} to {destfile} : {e}")
                    continue
                    #raise e
                    
            duration = (time.time()-start)/60.0
            print( f"  ... DONE. completed in {duration:1.3f} minutes")
            rv.append (destfile)
            
            if '.zip' in fname:
                print(f"unzipping {destfile} ..." )
                
                with zipfile.ZipFile(destfile, 'r') as zip_ref:
                    zip_ref.extractall(os.path.dirname(destfile))  
                
                log.info(f'{fname} downloaded and unzipped.')
            else:
                log.info(f'{fname} downloaded.')
                
    # for basename in ('wells.', 'unloc_wells.', 'cwilocs.zip'):
    #     for f in os.listdir(downloadpath):
    #         if f.startswith(basename):
    #             src = os.path.join(downloadpath, f)
    #             dst = os.path.join(downloadshppath, f)
    #             try:
    #                 shutil.move (src, dst)
    #             except Exception as e:
    #                 print (f'Moving "{src}" to "{dst}"')
    #                 print (e)
    
    duration = (time.time()-ftpstart)/60.0    
    print(f"*** ftp download & unzip finished in {duration:1.3f} minutes ***" )  
    
    if 0:
        # As of Sept 2025, observed no .xlsx files in downloaded files. 
        # This block of code is commented out to prevent accidental errors.
        # If .xlsx files exist, they may be local products rather than new data.
        # Run this block only after re-analysis!
        for dirpath, dirnames, filenames in os.walk(downloadpath):
            for file in filenames:
                if file.endswith(".xlsx"):
                    xlsxname = os.path.join(dirpath, file)
                    csvname = xlsxname.replace(".xlsx",".csv")
                    if os.path.exists(csvname):
                        print(f"Found csv:  {csvname}")  
                        with open(csvname, 'r') as file:
                            for i, line in enumerate(file):
                                print(line.strip())
                                if i > 6:  # Stop after 7 lines
                                    break
                        print ()
    
                    else:
                        print(f"Found xlsx: {xlsxname}  Converting ...", end=" ")
                        ok = xlsx_to_csv(xlsxname, csvname)
                        assert ok, csvname
                        print (csvname)
                           
    return rv

class ftp_params:
    '''
    The ftp login for the CWI download.

    All ftp login code is kept in this class.
    No credentials are written in this file.
    Method get() looks for the login in this order:
        1. File OWI_secrets.py (git-ignored; may be missing)
        2. The keyring
        3. Environment variables, named as in ftp_params.names
        4. Asking the user, who may then save it to the keyring or to an environment file
    Messages never display the password.
    '''
    names = ('CWI_FTP_SITE', 'CWI_FTP_DATA_PATH', 'CWI_FTP_USER', 'CWI_FTP_PASS')
    env_file = os.path.expanduser('~/.config/OWI/ftp.env')
    addr = None
    path = None
    user = None
    passwd = None

    def set_values(self, values):
        # Set all four values, or none of them.
        if values is None or len(values) != 4 or not all(values):
            return False
        self.addr, self.path, self.user, self.passwd = values
        return True

    def keyring_id(self):
        # The keyring service and username may be set in OWI_secrets.py
        service = getattr(S, 'CWI_FTP_KEYRING_SERVICE', '') or 'OWI_ftp'
        username = getattr(S, 'CWI_FTP_KEYRING_USER', '') or getpass.getuser()
        return service, username

    def get_from_secrets(self):
        values = [getattr(S, name, '') for name in self.names]
        return self.set_values(values)

    def get_from_keyring(self):
        # The keyring entry is stored as 'site|path|user|password'.
        # The password is last, so it may itself contain '|'.
        try:
            import keyring
            data = keyring.get_password(*self.keyring_id())
        except Exception as e:
            log.warning(f"Could not read the keyring ({type(e).__name__}).")
            return False
        if data is None:
            return False
        return self.set_values(data.split('|', 3))

    def get_from_environment(self):
        values = [os.getenv(name, '') for name in self.names]
        return self.set_values(values)

    def get_from_user(self):
        addr = input("Enter the FTP address: ")
        path = input(f"Enter the FTP path:  {addr}/")
        user = input("Enter your FTP username: ")
        passwd = getpass.getpass(prompt=f'Enter password for ftp user "{user}": ')
        return self.set_values((addr, path, user, passwd))

    def save_to_keyring(self):
        import keyring
        data = '|'.join((self.addr, self.path, self.user, self.passwd))
        keyring.set_password(*self.keyring_id(), data)
        print(f"FTP login for '{self.user}@{self.addr}' saved in the keyring.")

    def save_to_env_file(self, path=None):
        if path is None:
            path = self.env_file
        os.makedirs(os.path.dirname(path), exist_ok=True)
        values = (self.addr, self.path, self.user, self.passwd)
        # Make the file private before writing the password.
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        os.fchmod(fd, 0o600)
        with os.fdopen(fd, 'w') as f:
            for name, value in zip(self.names, values):
                f.write(f"export {name}={shlex.quote(value)}\n")
        print(f"FTP login saved in {path}")
        print("To load it in every new shell, add this line to ~/.bashrc:")
        print(f"    source {path}")

    def ask_to_save(self):
        answer = input("Save this FTP login for future runs? [k] keyring, [e] environment file, [n] no: ")
        answer = answer.strip().lower()
        if answer == 'k':
            self.save_to_keyring()
        elif answer == 'e':
            self.save_to_env_file()
        else:
            print("FTP login not saved.")

    def setup(self):
        # Ask the user for the login, and offer to save it.
        if not self.get_from_user():
            raise ValueError("FTP login is incomplete. All four values are required.")
        self.ask_to_save()

    def get(self):
        if self.get_from_secrets():
            log.info("FTP login found in OWI_secrets.py")
        elif self.get_from_keyring():
            log.info("FTP login found in the keyring")
        elif self.get_from_environment():
            log.info("FTP login found in environment variables")
        else:
            log.warning("FTP login not found in OWI_secrets.py, the keyring, or environment variables.")
            self.setup()
        return

def RUN_download_cwi(C):
    params = ftp_params()
    params.get()

    logging.debug("RUN_download_cwi()")
    logging.debug(versioninfo(C))
    
    logging.debug(f"OWI_DOWNLOAD_DIR = {C.OWI_DOWNLOAD_DIR}")
    logging.debug(f"OWI_DOWNLOAD_WELLSSHP_DIR = {C.OWI_DOWNLOAD_WELLSSHP_DIR}")
    logging.debug(f"OWI_DOWNLOAD_FILES = {C.OWI_DOWNLOAD_FILES}")
           
    newfiles = download_cwi_from_ftp(
        ftpaddr = params.addr,
        ftppath = params.path,
        ftpuser = params.user, 
        ftppass = params.passwd,
        downloadpath = C.OWI_DOWNLOAD_DIR,
        downloadshppath = C.OWI_DOWNLOAD_WELLSSHP_DIR,
        downloadfiles = C.OWI_DOWNLOAD_FILES) 
    print ('New files created:\n   ' + '\n   '.join(newfiles))
    return newfiles


if __name__ == '__main__':
    from OWI_config import OWI_version_40 as C
    if 0:   # Enter and save the ftp login only. No download.
        ftp_params().setup()
    if 1:
        RUN_download_cwi(C)
    
    print ('\n///////////// DONE ///////////////')